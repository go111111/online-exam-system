from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "在线考试系统课程设计报告-按要求重写版.docx"
ASSET_DIR = ROOT / "reports" / "assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)


def font(size=24, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for item in candidates:
        if Path(item).exists():
            return ImageFont.truetype(item, size=size)
    return ImageFont.load_default()


def text_size(draw, text, fnt):
    box = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=5, align="center")
    return box[2] - box[0], box[3] - box[1]


def draw_box(draw, xy, text, fill="#F7F9FC", outline="#2E74B5", width=3, radius=18, font_size=24, bold=False):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    fnt = font(font_size, bold=bold)
    tw, th = text_size(draw, text, fnt)
    draw.multiline_text((x1 + (x2 - x1 - tw) / 2, y1 + (y2 - y1 - th) / 2), text, fill="#222222", font=fnt, spacing=5, align="center")


def draw_arrow(draw, start, end, fill="#555555", width=3):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length = (dx * dx + dy * dy) ** 0.5 or 1
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 13
    p1 = (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55)
    p2 = (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55)
    draw.polygon([end, p1, p2], fill=fill)


def save_img(img, name):
    path = ASSET_DIR / name
    img.save(path)
    return path


def create_use_case_diagram():
    img = Image.new("RGB", (1600, 1050), "white")
    draw = ImageDraw.Draw(img)
    title_font = font(36, True)
    actor_font = font(24, True)
    use_case_font = font(21)
    draw.text((800, 45), "图1 在线考试系统用例图", fill="#111111", font=title_font, anchor="ma")

    # 系统边界。用例全部放在边界内部，演员放在边界外部，符合 UML 用例图习惯。
    boundary = (300, 110, 1300, 965)
    draw.rounded_rectangle(boundary, radius=18, outline="#1F4D78", width=4, fill="#FBFCFE")
    draw.text((800, 145), "在线考试系统", fill="#1F4D78", font=font(26, True), anchor="ma")

    def actor(cx, cy, label):
      draw.ellipse((cx - 28, cy - 92, cx + 28, cy - 36), outline="#333333", width=4, fill="#FFFFFF")
      draw.line((cx, cy - 36, cx, cy + 42), fill="#333333", width=4)
      draw.line((cx - 58, cy - 5, cx + 58, cy - 5), fill="#333333", width=4)
      draw.line((cx, cy + 42, cx - 50, cy + 105), fill="#333333", width=4)
      draw.line((cx, cy + 42, cx + 50, cy + 105), fill="#333333", width=4)
      draw.text((cx, cy + 140), label, fill="#111111", font=actor_font, anchor="ma")

    def use_case(cx, cy, text, fill="#FFFFFF"):
      rx, ry = 150, 48
      draw.ellipse((cx - rx, cy - ry, cx + rx, cy + ry), outline="#2E74B5", width=4, fill=fill)
      tw, th = text_size(draw, text, use_case_font)
      draw.multiline_text((cx - tw / 2, cy - th / 2), text, fill="#222222", font=use_case_font, spacing=4, align="center")
      return (cx - rx, cy, cx + rx, cy)

    def relation(start, end):
      draw.line((start, end), fill="#555555", width=3)

    actor(150, 405, "学生")
    actor(1450, 405, "管理员")

    student_cases = [
        ((520, 235), "查看考试列表"),
        ((520, 355), "参加在线考试"),
        ((520, 475), "保存/提交答案"),
        ((520, 595), "查看个人成绩"),
        ((520, 715), "接收系统通知"),
    ]
    admin_cases = [
        ((1080, 215), "考试管理"),
        ((1080, 335), "题库管理"),
        ((1080, 455), "试卷批改"),
        ((1080, 575), "成绩管理与统计"),
        ((1080, 695), "通知发布"),
        ((1080, 815), "用户维护"),
    ]

    # 学生只连接左列用例，管理员只连接右列用例。
    # 使用左右两条“关联总线”避免扇形交叉线，让版面更接近专业 UML 绘图工具输出。
    student_bus_x = 330
    admin_bus_x = 1270
    draw.line((student_bus_x, 235, student_bus_x, 715), fill="#555555", width=3)
    draw.line((admin_bus_x, 215, admin_bus_x, 815), fill="#555555", width=3)
    draw.line((210, 405, student_bus_x, 405), fill="#555555", width=3)
    draw.line((1390, 405, admin_bus_x, 405), fill="#555555", width=3)

    for (cx, cy), text in student_cases:
      left, mid_y, _, _ = use_case(cx, cy, text, fill="#E8EEF5")
      relation((student_bus_x, mid_y), (left, mid_y))

    for (cx, cy), text in admin_cases:
      _, mid_y, right, _ = use_case(cx, cy, text, fill="#F7F9FC")
      relation((right, mid_y), (admin_bus_x, mid_y))

    return save_img(img, "use_case.png")


def create_arch_diagram():
    img = Image.new("RGB", (1500, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((750, 45), "图2 系统总体架构设计图", fill="#111111", font=font(34, True), anchor="ma")
    draw_box(draw, (110, 170, 420, 290), "前端表示层\nVue3 + TS + Router", fill="#E8EEF5", font_size=24)
    draw_box(draw, (595, 170, 905, 290), "接口层\nAxios + JWT", fill="#F2F4F7", outline="#1F4D78", font_size=24)
    draw_box(draw, (1080, 170, 1390, 290), "后端服务层\nExpress + REST API", fill="#E8EEF5", font_size=24)
    draw_box(draw, (315, 515, 675, 650), "业务模块\n考试/题库/作答/批改/通知/统计", fill="#FFFFFF", font_size=22)
    draw_box(draw, (850, 515, 1210, 650), "数据持久层\nSQLite/MySQL + uploads", fill="#FFFFFF", font_size=22)
    draw_arrow(draw, (420, 230), (595, 230))
    draw_arrow(draw, (905, 230), (1080, 230))
    draw_arrow(draw, (1235, 290), (1030, 515))
    draw_arrow(draw, (850, 585), (675, 585))
    draw_arrow(draw, (500, 515), (265, 290))
    draw.text((750, 760), "设计要点：前端负责交互，后端集中鉴权、评分与统计，数据库保存考试业务数据。", fill="#333333", font=font(22), anchor="ma")
    return save_img(img, "architecture.png")


def create_process_diagram():
    img = Image.new("RGB", (1500, 940), "white")
    draw = ImageDraw.Draw(img)
    draw.text((750, 45), "图3 在线考试核心流程设计图", fill="#111111", font=font(34, True), anchor="ma")
    steps = [
        ((60, 160, 255, 255), "管理员\n创建考试"),
        ((350, 160, 545, 255), "维护题目\n设置分值"),
        ((640, 160, 835, 255), "发布考试\n学生进入"),
        ((930, 160, 1125, 255), "学生作答\n保存答案"),
        ((1220, 160, 1415, 255), "提交试卷\n生成提交记录"),
        ((350, 550, 545, 645), "客观题\n自动评分"),
        ((640, 550, 835, 645), "主观题\n人工批改"),
        ((930, 550, 1125, 645), "刷新总分\n更新状态"),
        ((1220, 550, 1415, 645), "成绩查询\n统计分析"),
    ]
    for xy, text in steps:
        draw_box(draw, xy, text, fill="#F7F9FC")
    for start, end in [((255, 207), (350, 207)), ((545, 207), (640, 207)), ((835, 207), (930, 207)), ((1125, 207), (1220, 207))]:
        draw_arrow(draw, start, end)
    draw_arrow(draw, (1317, 255), (450, 550))
    draw_arrow(draw, (1317, 255), (740, 550))
    draw_arrow(draw, (545, 598), (640, 598))
    draw_arrow(draw, (835, 598), (930, 598))
    draw_arrow(draw, (1125, 598), (1220, 598))
    draw.text((750, 820), "本报告重点说明考试、题库、提交、评分、批改、统计等核心过程；登录只作为系统前置约束，不展开基础实现。", fill="#333333", font=font(20), anchor="ma")
    return save_img(img, "process.png")


def create_er_diagram():
    img = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((750, 45), "图4 数据库概念结构图", fill="#111111", font=font(34, True), anchor="ma")
    entities = [
        ((80, 160, 290, 245), "users\n用户"),
        ((430, 160, 640, 245), "exams\n考试"),
        ((780, 160, 990, 245), "questions\n题目"),
        ((1130, 160, 1410, 245), "question_options\n选项"),
        ((300, 500, 540, 585), "submissions\n提交记录"),
        ((650, 500, 860, 585), "answers\n答案"),
        ((980, 500, 1270, 585), "answer_files/\ndrawing_data"),
        ((650, 700, 860, 785), "notifications\n通知"),
    ]
    for xy, text in entities:
        draw_box(draw, xy, text, fill="#FFFFFF", font_size=22)
    links = [
        ((290, 202), (430, 202), "创建/关联"),
        ((640, 202), (780, 202), "包含"),
        ((990, 202), (1130, 202), "拥有"),
        ((185, 245), (420, 500), "提交"),
        ((535, 245), (420, 500), "对应"),
        ((540, 542), (650, 542), "包含"),
        ((860, 542), (980, 542), "附件/画布"),
        ((755, 700), (185, 245), "推送"),
    ]
    for start, end, label in links:
        draw_arrow(draw, start, end)
        draw.text(((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 - 8), label, fill="#555555", font=font(18), anchor="ma")
    return save_img(img, "er.png")


def set_run_font(run, font_name="宋体", size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def add_para(doc, text="", style=None, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color="2E74B5")
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color="2E74B5")
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(run, size=12, bold=True, color="1F4D78")
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, size=10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=10)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F7F9FC")
    set_cell_margins(cell, top=120, bottom=120, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9.5)
    doc.add_paragraph()


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.space_after = Pt(4)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("在线考试系统课程设计报告")
    set_run_font(run, size=9, color="666666")
    return doc


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("在线考试系统课程设计报告")
    set_run_font(r, size=24, bold=True, color="0B2545")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("按课程设计要求重写版")
    set_run_font(r2, size=14, color="1F4D78")
    doc.add_paragraph()
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("课题名称", "在线考试系统"),
            ("开发模式", "个人独立完成；使用 Vue 3 + TypeScript + Express + SQLite/MySQL"),
            ("报告重点", "需求分析、UML 建模、架构设计、核心流程、主要算法、调试问题、个人总结"),
            ("实现范围说明", "以当前项目实际实现为准；班级、专业、课程章节、随机组卷等作为扩展需求说明，不虚构为已完整实现。"),
            ("完成日期", "2026年7月2日"),
        ],
        widths=[3.0, 12.7],
    )
    doc.add_page_break()


def add_toc(doc):
    add_heading(doc, "目录", 1)
    for item in [
        "第一章 绪论",
        "第二章 系统需求分析",
        "第三章 系统概要设计",
        "第四章 系统详细设计",
        "第五章 系统实现",
        "第六章 调试过程及问题解决",
        "第七章 个人体会及总结",
        "参考文献",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()


def build():
    use_case = create_use_case_diagram()
    arch = create_arch_diagram()
    process = create_process_diagram()
    er = create_er_diagram()

    doc = setup_doc()
    add_cover(doc)
    add_toc(doc)

    add_heading(doc, "第一章 绪论", 1)
    add_heading(doc, "1.1 课题背景", 2)
    add_para(doc, "在线考试系统是教学管理信息化中的典型应用。传统纸质考试需要人工印刷、发卷、收卷、阅卷和统计成绩，流程长且容易出现统计误差。随着 Web 技术和数据库技术成熟，将考试创建、题库管理、学生作答、自动评分、人工批改、成绩查询和通知发布整合到一个系统中，可以明显提高考试组织效率。")
    add_heading(doc, "1.2 课题意义", 2)
    add_bullet(doc, "对教师或管理员而言，可以集中维护考试、题目、通知和成绩，减少重复工作。")
    add_bullet(doc, "对学生而言，可以在浏览器中完成在线作答、提交和成绩查询。")
    add_bullet(doc, "对课程设计而言，该系统覆盖用户、考试、题库、提交、评分、统计等典型业务，适合体现软件工程分析、设计和实现过程。")
    add_heading(doc, "1.3 开发内容概述", 2)
    add_para(doc, "本项目使用 Vue 3、TypeScript、Vite、Element Plus、Express、JWT、bcryptjs、Multer、SQLite/MySQL 等技术完成在线考试系统。报告重点不展开登录注册等基础算法，而是围绕题库管理、考试过程、答案提交、自动评分、主观题批改、成绩统计和通知管理等核心过程展开。")

    add_heading(doc, "第二章 系统需求分析", 1)
    add_heading(doc, "2.1 课程设计要求理解", 2)
    add_para(doc, "根据课程设计要求，本系统需要体现综合应用所学知识完成一个软件系统的过程，报告中必须给出需求分析、UML 建模、概要设计、详细设计、系统实现、调试问题和个人总结。系统主题选择为在线考试系统，实际实现以 Web 端在线考试为主。")
    add_heading(doc, "2.2 功能需求分析", 2)
    add_table(
        doc,
        ["需求类别", "课程设计要求", "本项目实现情况"],
        [
            ("登录/退出", "登录、退出和基础权限控制", "已实现 JWT 登录态、管理员路由保护和退出。报告不重点展开基础算法。"),
            ("题目管理", "题目增删改查，题型可包括单选、多选、判断、分析、编程等", "已实现单选、填空、简答、绘图题；选择题选项独立存储。多选/判断/编程可作为后续扩展。"),
            ("题库/试卷管理", "组卷、随机选题、手动试卷等", "已实现考试创建、状态切换、题目维护。随机组卷和课程章节组卷作为扩展需求说明。"),
            ("用户管理", "教师、学生、班级、专业等", "已实现用户角色和管理员用户查询。班级、专业、课程归属可继续扩展。"),
            ("在线考试", "练习或考试，浏览器或移动端交互", "已实现浏览器在线考试、倒计时、离开页面提醒、自动交卷、答案保存。"),
            ("成绩管理", "成绩查询、批改、统计", "已实现客观题自动评分、主观题人工批改、学生成绩查询、管理员统计工作台。"),
            ("消息推送", "考试信息、成绩发布等主动通知", "已实现通知发布、通知列表、未读数量和已读标记。"),
            ("系统管理", "用户维护、权限和基础数据管理", "已实现管理员接口鉴权、用户列表、统计接口；基础数据管理仍可扩展。"),
        ],
        widths=[2.7, 5.6, 7.4],
    )
    add_heading(doc, "2.3 角色需求分析", 2)
    add_table(
        doc,
        ["角色", "目标", "主要操作"],
        [
            ("学生", "完成考试并查看结果", "查看开放考试、进入考试、保存答案、提交试卷、查看成绩、阅读通知。"),
            ("管理员", "组织考试并维护系统", "创建考试、管理题目、查看提交、批改主观题、打回重做、发布通知、查看统计。"),
            ("系统", "保证流程完整和数据一致", "鉴权、保存提交、计算客观题得分、汇总成绩、记录异常行为。"),
        ],
        widths=[2.5, 5.0, 8.2],
    )
    add_heading(doc, "2.4 UML 用例建模", 2)
    add_para(doc, "系统用例图如下。学生侧围绕参加考试和查看成绩展开，管理员侧围绕考试组织、题库维护、批改和统计展开。")
    doc.add_picture(str(use_case), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "2.5 非功能需求分析", 2)
    add_bullet(doc, "安全性：管理员接口必须经过 JWT 与角色校验，密码使用 bcrypt 加密。")
    add_bullet(doc, "可维护性：前端按页面、组件、类型文件组织；统计接口抽象出明确数据契约。")
    add_bullet(doc, "可靠性：提交、答案和分数分表保存，评分统一由后端刷新。")
    add_bullet(doc, "易用性：管理员首页提供统计卡片和近期提交，学生端提供倒计时、状态提示和成绩明细。")

    add_heading(doc, "第三章 系统概要设计", 1)
    add_heading(doc, "3.1 系统总体架构", 2)
    add_para(doc, "系统采用前后端分层架构：前端负责页面展示、路由跳转、表单和状态交互；后端提供 REST API、鉴权、评分、批改、通知和统计；数据库保存用户、考试、题目、提交、答案和通知数据。")
    doc.add_picture(str(arch), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "3.2 功能模块设计", 2)
    add_table(
        doc,
        ["模块", "核心页面/接口", "职责"],
        [
            ("考试管理模块", "AdminDashboard.vue、/api/admin/exams", "创建考试、切换状态、删除考试，并展示管理员统计面板。"),
            ("题目管理模块", "QuestionManager.vue、/api/admin/exams/:id/questions", "维护题目、选项、标准答案、分值和题目图片。"),
            ("在线考试模块", "ExamSession.vue、/api/exams/:id、/submit、/submit-answer", "加载试卷、倒计时、保存答案、上传附件或画布、提交试卷。"),
            ("评分批改模块", "AdminGrading.vue、refreshSubmissionScore", "客观题自动评分，主观题由管理员评分后刷新总分。"),
            ("成绩查询模块", "StudentResults.vue、AdminResults.vue、/api/user/results", "学生查看个人成绩，管理员查看所有提交。"),
            ("通知模块", "NotificationCenter.vue、AdminNotifications.vue", "通知发布、通知读取、未读计数和已读标记。"),
            ("统计分析模块", "adminAnalytics.ts、/api/admin/analytics/overview", "汇总考试数量、提交量、待批改数、平均得分率和通过率。"),
        ],
        widths=[2.8, 5.4, 7.5],
    )
    add_heading(doc, "3.3 核心业务流程设计", 2)
    doc.add_picture(str(process), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_heading(doc, "3.4 数据库概念设计", 2)
    add_para(doc, "数据库围绕用户、考试、题目、提交和答案建立关系。学生一次考试产生一条 submissions 记录，试卷中每一道题对应 answers 记录，主观题附件和绘图数据分别保存在 answer_files 与 drawing_data 中。")
    doc.add_picture(str(er), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(
        doc,
        ["表名", "作用", "关键字段"],
        [
            ("users", "用户与角色信息", "id、email、password、full_name、role"),
            ("exams", "考试基本信息", "id、title、start_time、end_time、duration_minutes、status"),
            ("questions", "题目主表", "id、exam_id、question_type、content、correct_answer、score、image_path"),
            ("question_options", "选择题选项", "question_id、option_label、option_text、is_correct"),
            ("submissions", "学生提交记录", "user_id、exam_id、status、submitted_at、total_score、cheated"),
            ("answers", "逐题答案与得分", "submission_id、question_id、student_answer、is_correct、awarded_score"),
            ("notifications", "系统通知", "title、content、type、target_role、created_by"),
        ],
        widths=[3, 5, 7.7],
    )

    add_heading(doc, "第四章 系统详细设计", 1)
    add_heading(doc, "4.1 题目管理详细设计", 2)
    add_para(doc, "题目管理是系统的基础模块。管理员选择考试后进入题目管理页面，可以新增、编辑、删除题目。页面中的题型被映射到后端数据库字段：choice 对应 single_choice，fill 对应 fill_blank，text 对应 short_answer，drawing 对应 drawing。")
    add_bullet(doc, "单选题：题干保存在 questions，选项保存在 question_options，正确答案用 is_correct 或 option_label 维护。")
    add_bullet(doc, "填空题：标准答案保存在 correct_answer，评分时做空白和大小写归一化。")
    add_bullet(doc, "简答题：提交后不自动判定正误，进入人工批改。")
    add_bullet(doc, "绘图题：通过画布或图片上传保存，管理员批改时查看附件或画布图片。")
    add_heading(doc, "4.2 在线考试过程详细设计", 2)
    add_para(doc, "在线考试页面加载考试信息和题目列表后，前端启动倒计时，并监听离开页面、切换标签、复制和右键操作。系统不把这些行为直接等同于作弊定论，而是记录异常标记并在达到阈值后自动交卷。")
    add_number(doc, "学生进入考试，前端请求 /api/exams/:id 获取考试和题目。")
    add_number(doc, "系统初始化倒计时，学生逐题作答。")
    add_number(doc, "文本答案保存在前端 answers 对象；附件和绘图题通过 /submit-answer 单题保存。")
    add_number(doc, "学生主动提交或倒计时结束后，调用 /api/exams/:id/submit。")
    add_number(doc, "后端保存 submissions 和 answers，然后调用评分函数刷新成绩状态。")
    add_heading(doc, "4.3 自动评分算法设计", 2)
    add_para(doc, "自动评分的关键不是直接比较原始字符串，而是先将学生答案和标准答案归一化。这样可以避免大小写、空格、多选顺序等因素导致误判。")
    add_code_block(
        doc,
        "算法：客观题自动评分\n"
        "输入：submissionId\n"
        "1. 查询该提交所属考试的所有题目和学生答案；\n"
        "2. 遍历题目：\n"
        "   2.1 如果是单选、填空等客观题，调用 normalizeAnswer 处理标准答案和学生答案；\n"
        "   2.2 两者一致则 awarded_score = question.score，否则 awarded_score = 0；\n"
        "   2.3 更新 answers.is_correct 和 answers.awarded_score；\n"
        "3. 如果发现简答题或绘图题未评分，则提交状态保持 submitted；\n"
        "4. 如果所有题目均可得出分数，则更新 submissions.total_score 和 status='graded'。"
    )
    add_heading(doc, "4.4 主观题批改算法设计", 2)
    add_para(doc, "主观题批改由管理员完成。系统先读取提交详情，再筛选 short_answer 和 drawing 类型题目。管理员输入每题得分后，后端校验分数不能小于 0 且不能超过题目分值，校验通过后写入 answers.awarded_score，并调用 refreshSubmissionScore 更新总分。")
    add_heading(doc, "4.5 统计分析算法设计", 2)
    add_para(doc, "管理员统计工作台通过 /api/admin/analytics/overview 聚合多张表数据。该接口先统计考试数、学生数、提交数，再按考试汇总题目总分和提交记录，最后计算平均得分率、通过率、待批改数量和异常提交数量。")
    add_code_block(
        doc,
        "算法：管理员统计分析\n"
        "1. 查询 exams、users、submissions 的总量；\n"
        "2. 按 exam_id 汇总 questions 的题量和满分；\n"
        "3. 查询已提交试卷并关联学生和考试名称；\n"
        "4. 对 status='graded' 且满分大于 0 的记录计算得分率；\n"
        "5. 得分率 >= 60% 计为通过；\n"
        "6. 返回 totals、examPerformance、latestSubmissions 三部分数据。"
    )

    add_heading(doc, "第五章 系统实现", 1)
    add_heading(doc, "5.1 前端核心实现", 2)
    add_para(doc, "前端入口在 src/main.ts 中创建 Vue 应用、注册 Element Plus、配置路由和 axios 拦截器。管理员端主要页面包括 AdminDashboard.vue、QuestionManager.vue、AdminGrading.vue、AdminResults.vue、AdminNotifications.vue；学生端主要页面包括 Dashboard.vue、ExamSession.vue、StudentResults.vue。")
    add_table(
        doc,
        ["页面/组件", "实现内容"],
        [
            ("AdminDashboard.vue", "考试管理和统计工作台，展示考试总数、学生数、提交数、平均得分率、考试表现和近期提交。"),
            ("QuestionManager.vue", "题目新增、编辑、删除、图片上传和选择题选项维护。"),
            ("ExamSession.vue", "考试作答、倒计时、异常行为记录、单题保存和最终提交。"),
            ("AnswerSubmitter.vue", "封装文本答案、附件答案和绘图题提交交互。"),
            ("AdminGrading.vue", "查看提交详情，为主观题评分，并同步刷新成绩。"),
            ("StudentResults.vue", "学生成绩概览和成绩明细。"),
        ],
        widths=[4.2, 11.5],
    )
    add_heading(doc, "5.2 后端核心实现", 2)
    add_para(doc, "后端集中在 server.ts 中实现。虽然该文件仍可进一步拆分，但当前已经包含数据库初始化、鉴权中间件、考试接口、题目接口、提交接口、批改接口、通知接口和统计接口。")
    add_bullet(doc, "数据库兼容设计：优先连接 MySQL，失败时回退到 SQLite，方便本地课程设计演示。")
    add_bullet(doc, "文件上传：使用 Multer 保存图片和答题附件，静态读取通过 /api/uploads/:filename 完成。")
    add_bullet(doc, "权限控制：管理员接口统一经过 authenticateToken 和 isAdmin。")
    add_bullet(doc, "评分核心：refreshSubmissionScore 作为统一成绩刷新入口，避免多个接口各自计算总分。")
    add_heading(doc, "5.3 核心过程实现标注", 2)
    add_table(
        doc,
        ["过程", "涉及接口/文件", "说明"],
        [
            ("题目维护过程", "QuestionManager.vue；/api/admin/exams/:id/questions", "管理员维护题干、题型、标准答案、分值和图片。"),
            ("考试作答过程", "ExamSession.vue；/api/exams/:id", "加载考试并初始化倒计时，学生作答过程中保存答案。"),
            ("提交评分过程", "/api/exams/:id/submit；refreshSubmissionScore", "保存提交记录，自动评分客观题，主观题保留待批改。"),
            ("主观题批改过程", "AdminGrading.vue；/api/admin/submissions/:id/grade", "管理员输入得分，后端校验分数范围并更新总分。"),
            ("打回重做过程", "/api/admin/submissions/:id/reject", "管理员打回后清理旧答案，学生可以重新进入考试。"),
            ("统计分析过程", "/api/admin/analytics/overview；adminAnalytics.ts", "后端聚合统计口径，前端按类型契约展示。"),
            ("通知处理过程", "NotificationCenter.vue；/api/notifications", "根据 target_role 推送通知并记录已读状态。"),
        ],
        widths=[3.2, 5.8, 6.7],
    )
    add_heading(doc, "5.4 与课程要求的差异说明", 2)
    add_para(doc, "课程要求中列举了教师、学生、班级、专业、课程章节、随机组卷、手机端 APP、判断题、多选题、分析题和编程题等功能。当前项目以在线考试核心闭环为主，已经实现学生/管理员两类角色、考试管理、题库管理、在线考试、批改、成绩查询、通知和统计；班级专业、课程章节、随机组卷、多选判断编程题等功能在本报告中作为后续扩展，不虚构为已完成内容。")

    add_heading(doc, "第六章 调试过程及问题解决", 1)
    add_heading(doc, "6.1 Vite 构建失败问题", 2)
    add_para(doc, "问题现象：执行 npm run build 时，Vite 报错 StudentResults.vue 至少需要一个 template 或 script。")
    add_para(doc, "原因分析：StudentResults.vue 被清空为 0 字节，Vue SFC 编译器无法解析。")
    add_para(doc, "解决方法：重建 StudentResults.vue，恢复学生成绩概览和成绩明细表，重新执行 npm run build 后通过。")
    add_heading(doc, "6.2 题目选项丢失问题", 2)
    add_para(doc, "问题现象：选择题如果只把选项写在前端状态中，刷新或重新进入题目管理页面后选项可能无法稳定恢复。")
    add_para(doc, "解决方法：后端使用 question_options 表保存选项，新增和编辑题目时同步写入，读取题目时再按 question_id 查询并返回给前端。")
    add_heading(doc, "6.3 主观题成绩状态问题", 2)
    add_para(doc, "问题现象：试卷既包含客观题又包含主观题时，如果提交后直接标记为已批改，会导致学生看到不完整成绩。")
    add_para(doc, "解决方法：refreshSubmissionScore 判断是否存在未评分主观题；如果存在，则 submission 状态保持 submitted，待管理员批改后再转为 graded。")
    add_heading(doc, "6.4 统计口径分散问题", 2)
    add_para(doc, "问题现象：如果前端分别请求考试、提交、题目列表再自行计算统计，会造成请求多、计算分散、口径不一致。")
    add_para(doc, "解决方法：新增 /api/admin/analytics/overview 后端聚合接口，统一返回 totals、examPerformance、latestSubmissions，前端只负责展示。")
    add_heading(doc, "6.5 测试结果", 2)
    add_table(
        doc,
        ["测试项", "命令/方式", "结果"],
        [
            ("类型检查", "npm run lint", "通过，tsc --noEmit 无报错。"),
            ("生产构建", "npm run build", "通过，生成 dist。"),
            ("健康检查", "GET /api/health", "返回 status=ok，数据库为 SQLite。"),
            ("统计接口", "管理员登录后访问 /api/admin/analytics/overview", "返回统计数据，包含 totals、examPerformance、latestSubmissions。"),
            ("源码打包", "排除 node_modules、dist、.env、数据库和上传文件后压缩", "得到干净源码包，适合提交。"),
        ],
        widths=[3, 6.6, 6.1],
    )

    add_heading(doc, "第七章 个人体会及总结", 1)
    add_para(doc, "本次课程设计让我更清楚地理解了一个系统不是把页面堆出来就结束，而是要从需求、角色、业务流程、数据库关系和异常情况出发设计。在线考试系统看起来只是“出题、考试、看成绩”，但真正实现时会遇到提交状态、客观题自动评分、主观题待批改、打回重做、附件保存、异常标记、通知已读等很多细节。")
    add_para(doc, "在实现过程中，我体会最深的是评分逻辑必须放在后端统一处理。如果前端直接计算成绩，不仅不安全，也容易造成不同页面的分数口径不一致。因此我把 refreshSubmissionScore 作为统一入口，让提交、批改和成绩查询都依赖同一套规则。")
    add_para(doc, "另一个收获是后台管理系统需要关注信息层级。最开始管理员首页只有考试列表，后来加入统计卡片、考试表现和近期提交后，管理员可以先看到系统整体运行情况，再进入具体考试或批改页面，这比单纯 CRUD 更接近真实管理系统。")
    add_para(doc, "不足之处是当前后端代码仍集中在 server.ts 中，后续如果继续完善，应拆分 routes、controllers、services 和 repositories；题型方面也可以继续扩展多选题、判断题、分析题、编程题；业务方面可以增加班级、专业、课程章节、随机组卷和错题分析等功能。")
    add_para(doc, "总体来说，本次课程设计完成了在线考试系统的核心闭环，也让我对 Vue 前端、Express 后端、数据库设计、接口设计和软件工程报告写作有了更完整的认识。")

    add_heading(doc, "参考文献", 1)
    for ref in [
        "[1] Vue.js 官方文档，Vue 3 Composition API 与组件开发。",
        "[2] Express.js 官方文档，路由、中间件与 REST API 设计。",
        "[3] TypeScript 官方文档，类型系统与工程化实践。",
        "[4] SQLite 官方文档，关系型数据库表设计与 SQL 查询。",
        "[5] JSON Web Token 标准文档，基于 Token 的身份认证机制。",
        "[6] Element Plus 官方文档，Vue 3 后台管理界面组件使用。",
    ]:
        add_para(doc, ref)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
