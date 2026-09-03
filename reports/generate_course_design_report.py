from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = "reports/在线考试系统课程设计报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_run_font(run, font_name="宋体", size=None, bold=None, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, align=None, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, bold=bold)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color="2E74B5")
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color="2E74B5")
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(run, size=12, bold=True, color="1F4D78")
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_width(table)
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_shading(hdr_cells[i], "F2F4F7")
        set_cell_margins(hdr_cells[i])
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, size=10.5, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=10)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("在线考试系统课程设计报告")
    set_run_font(footer_run, size=9, color="666666")
    return doc


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("在线考试系统课程设计报告")
    set_run_font(run, size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Vue 3 + Express + SQLite/MySQL 的在线考试管理平台")
    set_run_font(r, size=13, color="1F4D78")
    doc.add_paragraph()
    info = [
        ("项目名称", "在线考试系统"),
        ("技术栈", "Vue 3、TypeScript、Vite、Express、JWT、SQLite/MySQL"),
        ("本次新增", "管理员考试数据统计、考试表现分析、近期提交监控、学生成绩页恢复"),
        ("本次优化", "后端聚合统计接口、前端统计类型契约、管理员首页信息架构优化"),
        ("生成日期", "2026年7月1日"),
    ]
    table = add_table(doc, ["项目", "内容"], info, widths=[3.2, 12.5])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc.add_page_break()


def add_toc(doc):
    heading(doc, "目录", 1)
    chapters = [
        "第一章 绪论",
        "第二章 系统需求分析",
        "第三章 系统总体设计",
        "第四章 数据库设计",
        "第五章 系统详细设计与实现",
        "第六章 系统测试",
        "第七章 总结与展望",
        "参考文献",
        "致谢",
    ]
    for item in chapters:
        add_bullet(doc, item)
    doc.add_page_break()


def build_report():
    doc = setup_document()
    add_cover(doc)
    add_toc(doc)

    heading(doc, "第一章 绪论", 1)
    heading(doc, "1.1 课题背景", 2)
    paragraph(doc, "随着高校和培训机构教学过程逐步数字化，传统纸质考试在组织、阅卷、成绩统计和资料归档方面暴露出效率低、人工误差高、数据难复用等问题。在线考试系统可以把考试创建、题库维护、学生作答、自动评分、人工批改、成绩查询和通知发布串联为统一流程，减少人工流转成本。")
    paragraph(doc, "本项目面向课程设计场景，采用前后端分离思路：前端使用 Vue 3 与 TypeScript 构建交互页面，后端使用 Express 提供 REST API，数据层支持 SQLite 开发环境和 MySQL 部署环境。")
    heading(doc, "1.2 课题意义", 2)
    add_bullet(doc, "提升考试管理效率：管理员可以统一创建考试、维护题目、发布通知并查看成绩。")
    add_bullet(doc, "提升评分与统计能力：客观题自动评分，主观题保留人工批改入口，成绩数据可用于后续分析。")
    add_bullet(doc, "增强系统可扩展性：通过接口层、数据库表设计和前端组件化，把考试、通知、批改、成绩等模块拆分清楚。")
    heading(doc, "1.3 主要研究内容", 2)
    paragraph(doc, "本系统主要研究在线考试业务闭环的设计与实现，包括用户认证、题库管理、试卷管理、在线作答、提交与批改、成绩查询、通知管理和统计分析。")
    add_bullet(doc, "本次新增：管理员首页新增考试数据统计面板，包含考试总数、学生人数、提交试卷数、待批改数量、平均得分率、通过率、考试表现分析和近期提交监控。")
    add_bullet(doc, "本次新增：后端新增 /api/admin/analytics/overview 接口，统一聚合 exams、questions、submissions、users 数据。")
    add_bullet(doc, "本次优化：新增 src/types/adminAnalytics.ts 作为前端统计接口契约，减少页面层对 any 字段的盲目依赖。")
    add_bullet(doc, "本次优化：恢复并简化学生成绩页，使 /results 页面重新具备成绩概览和成绩明细展示能力，解决空组件导致构建失败的问题。")

    heading(doc, "第二章 系统需求分析", 1)
    heading(doc, "2.1 可行性分析", 2)
    add_table(doc, ["维度", "分析结果"], [
        ("技术可行性", "Vue 3、TypeScript、Express、JWT、SQLite/MySQL 均为成熟技术，适合课程设计和小型管理系统。"),
        ("经济可行性", "项目依赖开源技术和本地数据库，开发与运行成本低。"),
        ("操作可行性", "系统提供管理员和学生两类入口，核心操作以表单、表格和按钮完成，学习成本较低。"),
        ("维护可行性", "前端按 views/components/types 划分，后端接口集中在 server.ts，后续可继续拆分 service/controller。"),
    ], widths=[3, 12.7])
    heading(doc, "2.2 用户角色分析", 2)
    add_table(doc, ["角色", "主要职责", "关键权限"], [
        ("学生", "参加考试、提交答案、查看成绩和通知", "登录、查看开放考试、作答、提交、查看个人成绩"),
        ("管理员", "维护考试、题库、通知、批改和成绩", "创建考试、管理题目、发布通知、批改试卷、查看统计"),
    ], widths=[2.5, 6.5, 6.7])
    heading(doc, "2.3 功能需求分析", 2)
    add_bullet(doc, "登录注册：支持邮箱、密码和角色信息，使用 JWT 维护登录态。")
    add_bullet(doc, "考试管理：管理员可创建、开启、关闭和删除考试。")
    add_bullet(doc, "题库管理：支持选择题、填空题、简答题、绘图题，题目可携带图片。")
    add_bullet(doc, "在线考试：学生进入考试后作答，系统记录提交时间、答题用时、异常标记。")
    add_bullet(doc, "成绩管理：客观题自动评分，主观题人工批改，学生和管理员均可查看成绩。")
    add_bullet(doc, "通知管理：管理员发布公告、考试通知或系统通知，学生端可查看并标记已读。")
    add_bullet(doc, "统计分析：管理员查看整体考试运行情况、通过率、平均得分率和近期提交。")
    heading(doc, "2.4 非功能需求分析", 2)
    add_table(doc, ["需求类别", "要求"], [
        ("安全性", "使用 JWT 鉴权，管理员接口增加角色校验，密码使用 bcrypt 加密保存。"),
        ("易用性", "页面信息层级清晰，使用状态标签区分开启、待批改、已批改和打回。"),
        ("可靠性", "提交记录与答案分表存储，成绩刷新逻辑统一计算总分。"),
        ("可维护性", "新增统计接口返回稳定 JSON 结构，前端使用 TypeScript interface 描述数据契约。"),
        ("可扩展性", "数据库同时考虑 SQLite 和 MySQL，后续可扩展班级、课程、组卷策略等模块。"),
    ], widths=[3, 12.7])

    heading(doc, "第三章 系统总体设计", 1)
    heading(doc, "3.1 系统架构设计", 2)
    paragraph(doc, "系统采用浏览器客户端、Express API 服务和数据库三层结构。前端负责视图渲染、路由控制、表单交互和接口调用；后端负责身份认证、权限判断、业务规则、评分逻辑和文件上传；数据库保存用户、考试、题目、答案、成绩和通知等业务数据。")
    add_table(doc, ["层次", "技术/文件", "职责"], [
        ("前端表示层", "src/views、src/components、Vue Router", "展示学生端和管理员端页面，处理表单、表格、状态标签和操作按钮。"),
        ("前端接口契约", "src/types/adminAnalytics.ts", "定义统计接口返回结构，保证页面使用字段明确。"),
        ("后端服务层", "server.ts、Express、JWT", "提供 REST API、鉴权、角色校验、评分、通知、统计聚合。"),
        ("数据持久层", "SQLite/MySQL", "保存 users、exams、questions、submissions、answers 等数据。"),
        ("文件存储", "uploads/", "保存题目图片和答题附件。"),
    ], widths=[3, 4.5, 8.2])
    paragraph(doc, "架构流程：用户浏览器 -> Vue 路由页面 -> axios API 调用 -> Express 鉴权与业务处理 -> SQLite/MySQL 查询 -> JSON 返回 -> 页面更新。")
    heading(doc, "3.2 系统功能模块设计", 2)
    add_table(doc, ["模块", "子功能", "实现位置"], [
        ("用户认证模块", "注册、登录、JWT 鉴权、角色判断", "server.ts、Login.vue、Register.vue、main.ts"),
        ("考试管理模块", "考试创建、状态切换、删除、考试列表", "AdminDashboard.vue、server.ts"),
        ("题库管理模块", "题目增删改、选项管理、图片上传", "QuestionManager.vue、server.ts"),
        ("在线考试模块", "考试详情、答题、单题保存、最终提交", "ExamSession.vue、server.ts"),
        ("批改管理模块", "主观题评分、打回重做、最终成绩刷新", "AdminGrading.vue、AdminResults.vue、server.ts"),
        ("通知模块", "通知发布、通知列表、未读数量、标记已读", "AdminNotifications.vue、NotificationCenter.vue、server.ts"),
        ("统计分析模块", "管理员总览、考试表现、近期提交", "AdminDashboard.vue、adminAnalytics.ts、server.ts"),
    ], widths=[3, 7, 5.7])
    heading(doc, "3.3 系统流程设计", 2)
    add_number(doc, "学生登录系统，系统校验 JWT 后进入考试列表。")
    add_number(doc, "学生选择开放考试，系统读取考试详情和题目列表。")
    add_number(doc, "学生作答并提交，后端保存 submissions 与 answers，并自动评分客观题。")
    add_number(doc, "若试卷含主观题，管理员进入批改页面给出主观题得分。")
    add_number(doc, "系统刷新总分、状态和批改时间，学生端与管理员端均可查看结果。")
    add_number(doc, "管理员首页通过统计接口聚合运行数据，辅助判断考试完成情况和批改压力。")

    heading(doc, "第四章 数据库设计", 1)
    heading(doc, "4.1 数据库概念设计", 2)
    paragraph(doc, "系统核心实体包括用户、考试、题目、题目选项、提交记录、答案、答案附件、绘图数据、通知和通知已读记录。用户与提交记录是一对多关系，考试与题目是一对多关系，考试与提交记录是一对多关系，提交记录与答案是一对多关系，通知与已读记录是一对多关系。")
    heading(doc, "4.2 数据库表设计", 2)
    add_table(doc, ["表名", "主要字段", "说明"], [
        ("users", "id、email、password、full_name、role", "保存用户账号、加密密码和角色。"),
        ("exams", "id、title、description、start_time、end_time、duration_minutes、status、created_by", "保存考试基础信息和发布状态。"),
        ("questions", "id、exam_id、question_type、content、correct_answer、score、sort_order、image_path", "保存题目内容、题型、答案和分值。"),
        ("question_options", "id、question_id、option_label、option_text、is_correct", "保存选择题选项。"),
        ("submissions", "id、user_id、exam_id、status、submitted_at、total_score、used_time_minutes、cheated", "保存学生提交记录和试卷状态。"),
        ("answers", "id、submission_id、question_id、student_answer、is_correct、awarded_score、submission_type", "保存每道题的学生答案和得分。"),
        ("answer_files", "id、answer_id、file_name、file_path、file_size、file_type", "保存答题附件元数据。"),
        ("drawing_data", "id、answer_id、canvas_json", "保存绘图题画布数据。"),
        ("notifications", "id、title、content、type、target_role、exam_id、created_by、created_at", "保存管理员发布的通知。"),
        ("notification_reads", "id、notification_id、user_id、read_at", "保存用户通知已读状态。"),
    ], widths=[3.4, 7.8, 4.5])

    heading(doc, "第五章 系统详细设计与实现", 1)
    heading(doc, "5.1 登录功能实现", 2)
    paragraph(doc, "登录接口接收邮箱和密码，后端查询用户并使用 bcrypt 比对密码，验证通过后生成 JWT。前端在 main.ts 中维护 auth 对象，把 token 写入 localStorage，并通过 axios 拦截器自动附带 Authorization 请求头。")
    heading(doc, "5.2 用户管理功能实现", 2)
    paragraph(doc, "用户表通过 role 区分 student 和 admin。管理员接口统一使用 authenticateToken 与 isAdmin 中间件保护，避免普通学生调用考试管理、批改和统计接口。")
    heading(doc, "5.3 题库管理功能实现", 2)
    paragraph(doc, "题库支持选择题、填空题、简答题和绘图题。选择题除 questions 表保存题干和答案外，还通过 question_options 表保存选项文本和正确标记；图片题通过 Multer 上传到 uploads 目录，并在题目记录中保存 image_path。")
    heading(doc, "5.4 试卷管理功能实现", 2)
    paragraph(doc, "管理员在 AdminDashboard.vue 创建考试，填写标题、描述、开始时间、结束时间、考试时长和状态。后端写入 exams 表，管理员可通过状态按钮将考试切换为 open 或 closed。")
    heading(doc, "5.5 在线考试功能实现", 2)
    paragraph(doc, "学生从 Dashboard.vue 进入考试，ExamSession.vue 加载题目并记录作答。提交后，后端创建或更新 submissions，并逐题保存 answers。客观题通过 normalizeAnswer 统一格式后自动判分，主观题进入待批改状态。")
    heading(doc, "5.6 成绩管理功能实现", 2)
    paragraph(doc, "系统通过 refreshSubmissionScore 统一计算试卷总分：选择题、填空题等客观题自动给分，简答题和绘图题等待管理员人工给分。管理员可在 AdminGrading.vue 批改，AdminResults.vue 查看所有成绩和异常标记，学生可在 StudentResults.vue 查看个人成绩。")
    heading(doc, "5.7 本次新增：管理员统计分析功能", 2)
    paragraph(doc, "本次新增 /api/admin/analytics/overview 接口，后端聚合考试数量、学生数量、提交数量、待批改数量、异常数量、平均得分率、通过率、单场考试表现和近期提交记录。前端 AdminDashboard.vue 在考试列表上方新增四个统计卡片、考试表现分析表和近期提交列表。")
    add_table(doc, ["新增/优化点", "对应文件", "说明"], [
        ("新增统计接口", "server.ts", "新增 /api/admin/analytics/overview，统一聚合多个表的数据。"),
        ("新增统计类型契约", "src/types/adminAnalytics.ts", "定义 totals、examPerformance、latestSubmissions 的字段类型。"),
        ("优化管理员首页", "src/views/AdminDashboard.vue", "加入统计卡片、考试表现分析、近期提交监控，并在考试增删改后同步刷新统计。"),
        ("修复学生成绩页", "src/views/StudentResults.vue", "从空组件恢复为可用的成绩概览和明细表，解决 Vite 构建失败。"),
        ("新增数值处理工具", "server.ts", "新增 toNumber 与 roundToOne，保证统计计算对空值更稳定。"),
    ], widths=[3.5, 5.2, 7])
    heading(doc, "5.8 本次架构优化说明", 2)
    add_bullet(doc, "接口聚合前移到后端：统计口径集中在服务端，前端不再重复拉取多个列表自行计算。")
    add_bullet(doc, "数据契约显式化：新增 TypeScript interface，使统计页面依赖的数据结构可读、可维护。")
    add_bullet(doc, "管理员首页从单一考试列表升级为工作台：先看整体数据和风险，再进入具体考试或批改页面。")
    add_bullet(doc, "成绩页恢复可构建状态：避免空 Vue SFC 阻断生产构建，保证课程设计演示链路完整。")

    heading(doc, "第六章 系统测试", 1)
    heading(doc, "6.1 测试目的", 2)
    paragraph(doc, "测试目标是确认系统主要功能链路可用，新增统计功能能够通过接口返回真实数据，前端项目能够通过类型检查和生产构建。")
    heading(doc, "6.2 测试环境", 2)
    add_table(doc, ["项目", "环境"], [
        ("操作系统", "Windows，项目目录 E:\\code\\Project\\online-exam-system"),
        ("前端构建", "Vite、Vue 3、TypeScript"),
        ("后端运行", "npm run dev，Express 服务监听 http://localhost:3000"),
        ("数据库", "本地 SQLite exam.db，MySQL 不可用时自动回退"),
        ("测试日期", "2026年7月1日"),
    ], widths=[3.2, 12.5])
    heading(doc, "6.3 功能测试", 2)
    add_table(doc, ["测试项", "测试方法", "结果"], [
        ("类型检查", "执行 npm run lint", "通过，tsc --noEmit 无报错。"),
        ("生产构建", "执行 npm run build", "通过，dist 构建成功；仅提示部分 chunk 体积较大。"),
        ("健康检查", "访问 GET /api/health", "返回 status=ok，database=SQLite。"),
        ("管理员登录", "POST /api/login，使用默认管理员账号", "返回 token，可访问管理员接口。"),
        ("统计接口", "GET /api/admin/analytics/overview", "返回 totals、examPerformance、latestSubmissions 数据。"),
        ("成绩页修复", "Vite build 重新解析 StudentResults.vue", "通过，空组件错误消失。"),
    ], widths=[3.2, 7, 5.5])

    heading(doc, "第七章 总结与展望", 1)
    heading(doc, "7.1 总结", 2)
    paragraph(doc, "本课程设计完成了在线考试系统的主要业务闭环：用户登录、考试管理、题库维护、学生在线作答、提交记录保存、自动评分、人工批改、成绩查询、通知管理和管理员统计分析。系统采用 Vue 3 + Express + SQLite/MySQL 的技术组合，能够满足课程设计对功能完整性、可运行性和可说明性的要求。")
    paragraph(doc, "本次优化重点放在管理员工作台和统计能力上，通过新增统计接口和类型契约，把原本分散在各表中的数据整理为面向管理决策的总览指标，使系统不只能够完成考试流程，也能反映考试运行状态。")
    heading(doc, "7.2 不足与展望", 2)
    add_bullet(doc, "后端目前仍集中在 server.ts，后续可拆分为 routes、controllers、services、repositories，提高长期维护性。")
    add_bullet(doc, "统计功能当前以基础指标为主，后续可增加按班级、课程、题型、知识点维度的分析。")
    add_bullet(doc, "考试防作弊能力仍较基础，后续可加入切屏记录、摄像头校验、答题行为日志等机制。")
    add_bullet(doc, "题库可继续扩展批量导入、随机组卷、错题分析和难度系数统计。")
    add_bullet(doc, "部署层面可补充 Docker Compose、Nginx 和 MySQL 的生产环境配置文档。")

    heading(doc, "参考文献", 1)
    refs = [
        "[1] Vue.js 官方文档：Vue 3 Composition API 与组件开发。",
        "[2] Express.js 官方文档：路由、中间件与 REST API 设计。",
        "[3] TypeScript 官方文档：接口类型、类型检查与工程化实践。",
        "[4] SQLite 官方文档：轻量级关系型数据库设计与 SQL 查询。",
        "[5] JSON Web Token 标准：基于 Token 的身份认证机制。",
    ]
    for ref in refs:
        paragraph(doc, ref)

    heading(doc, "致谢", 1)
    paragraph(doc, "感谢课程设计过程中提供的需求方向和项目基础代码。通过本次功能增强与报告整理，系统在考试管理、成绩统计和课程设计说明方面更加完整，也为后续继续扩展在线教学平台功能奠定了基础。")

    doc.save(OUT)


if __name__ == "__main__":
    build_report()
